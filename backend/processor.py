import os
from dotenv import load_dotenv
import PyPDF2
import re
from duckduckgo_search import DDGS
from docx import Document
from sentence_transformers import SentenceTransformer
from backend.database import chunks_collection, documents_collection, notes_collection, quizzes_collection
from bson import ObjectId
from datetime import datetime
from groq import Groq

load_dotenv()

# Load models once
model = SentenceTransformer('all-MiniLM-L6-v2')
groq_client = Groq(api_key=os.getenv("GROQ_API_KEY"))

def process_document(file_path, filename, subject, owner_email, chat_id=None):
    """
    Processes a document and indexes it. 
    If chat_id is provided, the chunks are marked as temporary for that specific chat.
    """
    text = ""
    try:
        if filename.endswith(".pdf"):
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        elif filename.endswith(".docx"):
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()

        print(f"Extracted {len(text)} characters from {filename}")
        if not text.strip():
            print(f"No text extracted from {filename}")
            return None

        # Create Document Entry
        doc_entry = {
            "filename": filename,
            "subject": subject,
            "owner": owner_email,
            "upload_date": datetime.utcnow(),
            "chat_id": chat_id, # If None, it's global (Faculty)
            "is_temporary": chat_id is not None
        }
        doc_id = documents_collection.insert_one(doc_entry).inserted_id

        # Chunking
        chunks = []
        chunk_size = 500
        for i in range(0, len(text), chunk_size):
            chunk_text = text[i:i+chunk_size]
            embedding = model.encode(chunk_text).tolist()
            chunks.append({
                "document_id": doc_id,
                "text": chunk_text,
                "embedding": embedding,
                "chat_id": chat_id, # Tag chunks for temporary search
                "is_temporary": chat_id is not None
            })
        
        if chunks:
            chunks_collection.insert_many(chunks)
        
        return doc_id, text
    except Exception as e:
        print(f"Processing Error: {e}")
        return None

def query_system(question, chat_id=None):
    """
    Search both global content AND chat-specific content.
    Filters in Python to avoid MongoDB index requirement on chat_id.
    """
    try:
        query_vector = model.encode(question).tolist()
        
        pipeline = [
            {
                "$vectorSearch": {
                    "index": "vector_index",
                    "path": "embedding",
                    "queryVector": query_vector,
                    "numCandidates": 100,
                    "limit": 50
                }
            },
            {
                "$project": {
                    "text": 1,
                    "score": {"$meta": "searchScore"},
                    "document_id": 1,
                    "chat_id": 1
                }
            }
        ]
        
        raw_results = list(chunks_collection.aggregate(pipeline))
        print(f"Vector search returned {len(raw_results)} raw results")
        
        # Filter manually for vector results
        results = []
        for r in raw_results:
            if r.get("chat_id") is None or r.get("chat_id") == chat_id:
                results.append(r)
        
        # STRONG CONSISTENCY FALLBACK:
        # Fetch the most recent chunks for this specific chat directly from the DB
        # This ensures that newly uploaded files are available even before Atlas Search indexes them.
        if chat_id:
            recent_chunks = list(chunks_collection.find({"chat_id": chat_id}).sort("_id", -1).limit(5))
            for rc in recent_chunks:
                # Avoid duplicates
                if not any(str(r.get("_id")) == str(rc.get("_id")) for r in results):
                    results.insert(0, rc) # Prepend to prioritize over global docs
                    
        print(f"Found {len(results)} relevant chunks (including recent uploads) for chat_id: {chat_id}")
        return results[:10] # Return top 10 combined
    except Exception as e:
        print(f"Query Error: {e}")
        return []

def generate_ai_response(question, context, history=None):
    """
    Generates an AI response using Groq based on context and history.
    """
    try:
        messages = [
            {"role": "system", "content": "You are EduNavigator AI, a professional academic assistant. Use the provided context to answer questions accurately. ALWAYS cite your sources by mentioning the filename (e.g., 'As per [filename]...') or using [filename] at the end of sentences. If the context is unhelpful, use your general knowledge but clearly state that it is from general knowledge and not the provided documents."}
        ]
        
        if history:
            for msg in history:
                messages.append({"role": msg["role"], "content": msg["content"]})
        
        messages.append({"role": "user", "content": f"Context: {context}\n\nQuestion: {question}"})
        
        response = groq_client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=messages,
            temperature=0.7,
            max_tokens=1024
        )
        
        return response.choices[0].message.content
    except Exception as e:
        print(f"AI Generation Error: {e}")
        return f"I'm sorry, I encountered an error while thinking: {str(e)}"

# --- NOTES GENERATION (MAP-REDUCE-FORMAT) ---

def clean_text(text):
    """Basic text cleaning: removing multiple spaces, headers/footers placeholders, etc."""
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'Page \d+ of \d+', '', text)
    return text.strip()

def map_summarize_chunks(text, chunk_size=4000):
    """Phase A: Chunk-level understanding - Summarize each chunk independently."""
    summaries = []
    for i in range(0, len(text), chunk_size):
        chunk = text[i : i + chunk_size]
        prompt = f"Summarize the following technical/academic content concisely, focusing on key facts and concepts:\n\n{chunk}"
        try:
            response = groq_client.chat.completions.create(
                model="llama-3.1-8b-instant", # Faster model for map phase
                messages=[{"role": "user", "content": prompt}],
                max_tokens=500
            )
            summaries.append(response.choices[0].message.content)
        except Exception as e:
            print(f"Chunk Summary Error: {e}")
    return summaries

def web_search_enrichment(topics):
    """Phase D: Web Search - Get external context for key topics."""
    web_results = []
    with DDGS() as ddgs:
        for topic in topics[:3]: # Limit to top 3 topics
            try:
                results = list(ddgs.text(topic, max_results=2))
                for r in results:
                    web_results.append({
                        "title": r['title'],
                        "href": r['href'],
                        "body": r['body']
                    })
            except Exception as e:
                print(f"Web Search Error for {topic}: {e}")
    return web_results

def generate_notes(doc_id, text, user_email):
    """Main orchestration for Map-Reduce-Format Notes Generation."""
    print(f"Starting notes generation for doc: {doc_id}")
    
    # 1. Clean Text
    text = clean_text(text)
    
    # 2. Map Phase
    chunk_summaries = map_summarize_chunks(text)
    combined_summaries = "\n\n".join(chunk_summaries)
    
    # 3. Identify Search Topics (Internal step)
    try:
        topic_prompt = f"Based on these document summaries, list the 3 most important academic topics or terms that would benefit from additional external context/verification. Return ONLY a comma-separated list.\n\nSummaries: {combined_summaries[:2000]}"
        topic_res = groq_client.chat.completions.create(
            model="llama-3.1-8b-instant",
            messages=[{"role": "user", "content": topic_prompt}]
        )
        topics = [t.strip() for t in topic_res.choices[0].message.content.split(",")]
    except:
        topics = []

    # 4. Web Search Phase
    web_context = web_search_enrichment(topics)
    web_text = "\n".join([f"Source: {w['href']}\nContent: {w['body']}" for w in web_context])

    # 5. Reduce & Format Phase
    final_prompt = f"""
    Using the provided Document Summaries and Web Search results, create a comprehensive set of structured notes.
    
    DOCUMENT SUMMARIES:
    {combined_summaries}
    
    WEB SEARCH CONTEXT:
    {web_text}
    
    STRUCTURE:
    1. SUMMARY: A high-level overview of the entire document.
    2. KEY POINTS: A bulleted list of the most important takeaways from the document.
    3. IMPORTANT CONCEPTS: Definitions or explanations of core concepts, enriched with web data where helpful.
    
    CITATION RULES:
    - Cite the document as [Document] for points originating from it.
    - Cite web sources using their URL, e.g., [https://example.com].
    - ALWAYS cite at least one source for every major point.
    
    Format the output as clean JSON with keys: "summary", "key_points", "important_concepts".
    Each key should contain a string or list of strings as appropriate.
    """
    
    try:
        response = groq_client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[{"role": "system", "content": "You are a professional academic note-taker. Output ONLY valid JSON."},
                      {"role": "user", "content": final_prompt}],
            response_format={"type": "json_object"}
        )
        notes_data = response.choices[0].message.content
        import json
        structured_notes = json.loads(notes_data)
        
        # Save to DB
        note_entry = {
            "document_id": ObjectId(doc_id),
            "user_email": user_email,
            "notes": structured_notes,
            "timestamp": datetime.utcnow()
        }
        notes_collection.update_one(
            {"document_id": ObjectId(doc_id)},
            {"$set": note_entry},
            upsert=True
        )
        print(f"Notes generated and saved for doc: {doc_id}")
        return structured_notes
    except Exception as e:
        print(f"Final Notes Generation Error: {e}")
        return None

# --- QUIZ GENERATION ---

def generate_quiz(doc_id, text, user_email):
    """Generates a structured practice quiz from document text."""
    print(f"Starting quiz generation for doc: {doc_id}")
    
    # Pick first few thousand characters to stay within token limits
    # In a more advanced version, we could pick diverse chunks.
    context = text[:8000]
    
    prompt = f"""
    Based on the following academic content, generate a practice quiz with 5 multiple-choice questions.
    
    CONTENT:
    {context}
    
    RULES:
    1. Each question must have exactly 4 options (A, B, C, D).
    2. Provide the index of the correct answer (0, 1, 2, or 3).
    3. Provide a brief explanation for why the answer is correct.
    4. Ensure the questions cover core concepts and facts from the text.
    
    Format the output as clean JSON with a key "questions" containing a list of objects.
    Each object should have keys: "question", "options" (list of 4 strings), "correct_index" (int), "explanation" (string).
    """
    
    try:
        response = groq_client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[{"role": "system", "content": "You are a professional academic tutor. Output ONLY valid JSON."},
                      {"role": "user", "content": prompt}],
            response_format={"type": "json_object"}
        )
        import json
        quiz_data = json.loads(response.choices[0].message.content)
        
        # Save to DB
        quiz_entry = {
            "document_id": ObjectId(doc_id),
            "user_email": user_email,
            "quiz": quiz_data,
            "timestamp": datetime.utcnow()
        }
        quizzes_collection.update_one(
            {"document_id": ObjectId(doc_id)},
            {"$set": quiz_entry},
            upsert=True
        )
        print(f"Quiz generated and saved for doc: {doc_id}")
        return quiz_data
    except Exception as e:
        print(f"Quiz Generation Error: {e}")
        return None
